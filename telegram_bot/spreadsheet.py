from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


EXPORT_COLUMNS = [
    ("sku", "SKU"),
    ("name", "Название"),
    ("category", "Категория"),
    ("price", "Цена продажи"),
    ("purchase_price", "Закупочная цена"),
    ("stock", "Остаток"),
    ("min_stock", "Минимальный остаток"),
    ("unit", "Ед. измерения"),
    ("supplier", "Поставщик"),
    ("barcode", "Штрихкод"),
    ("description", "Описание"),
    ("is_active", "Активен"),
]

ALIASES = {
    "sku": "sku", "артикул": "sku",
    "название": "name", "name": "name",
    "категория": "category", "category": "category",
    "цена": "price", "цена продажи": "price", "price": "price",
    "закупочная цена": "purchase_price", "purchase_price": "purchase_price",
    "остаток": "stock", "stock": "stock",
    "минимальный остаток": "min_stock", "минимум": "min_stock", "min_stock": "min_stock",
    "ед. измерения": "unit", "единица измерения": "unit", "unit": "unit",
    "поставщик": "supplier", "supplier": "supplier",
    "штрихкод": "barcode", "ean": "barcode", "barcode": "barcode",
    "описание": "description", "description": "description",
    "активен": "is_active", "is_active": "is_active",
}


def _normalize_header(value: Any) -> str | None:
    text = str(value or "").strip().casefold()
    return ALIASES.get(text)


def _normalize_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().casefold() not in {"0", "false", "нет", "no", "архив"}


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(row)
    normalized["sku"] = str(normalized.get("sku", "") or "").strip()
    normalized["name"] = str(normalized.get("name", "") or "").strip()
    normalized["category"] = str(normalized.get("category", "Без категории") or "Без категории").strip()
    for key in ("price", "purchase_price"):
        value = normalized.get(key, 0)
        normalized[key] = float(str(value or 0).replace(" ", "").replace(",", "."))
    for key in ("stock", "min_stock"):
        value = normalized.get(key, 0)
        normalized[key] = int(float(str(value or 0).replace(" ", "").replace(",", ".")))
    for key in ("unit", "supplier", "barcode", "description"):
        normalized[key] = str(normalized.get(key, "") or "").strip()
    normalized["unit"] = normalized["unit"] or "шт."
    normalized["is_active"] = _normalize_bool(normalized.get("is_active", True))
    return normalized


def export_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}1"

    for col, (_, title) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)

    for row_idx, item in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(EXPORT_COLUMNS, start=1):
            value = item.get(key, "")
            if key == "is_active":
                value = "Да" if value else "Нет"
            ws.cell(row=row_idx, column=col_idx, value=value)

    for idx, (_, title) in enumerate(EXPORT_COLUMNS, start=1):
        width = max(12, min(40, len(title) + 4))
        for cell in ws[get_column_letter(idx)]:
            if cell.value is not None:
                width = max(width, min(40, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_csv(rows: list[dict[str, Any]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow([title for _, title in EXPORT_COLUMNS])
    for item in rows:
        writer.writerow([
            ("Да" if item.get(key) else "Нет") if key == "is_active" else item.get(key, "")
            for key, _ in EXPORT_COLUMNS
        ])
    return out.getvalue().encode("utf-8-sig")


def parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    keys = [_normalize_header(value) for value in rows[0]]
    parsed: list[dict[str, Any]] = []
    for values in rows[1:]:
        raw = {key: value for key, value in zip(keys, values) if key}
        if not any(value not in (None, "") for value in raw.values()):
            continue
        parsed.append(_normalize_row(raw))
    return parsed


def parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return []
    keys = [_normalize_header(value) for value in rows[0]]
    parsed: list[dict[str, Any]] = []
    for values in rows[1:]:
        raw = {key: value for key, value in zip(keys, values) if key}
        if not any(str(value).strip() for value in raw.values()):
            continue
        parsed.append(_normalize_row(raw))
    return parsed


def parse_catalog_file(filename: str, data: bytes) -> list[dict[str, Any]]:
    lower = filename.casefold()
    if lower.endswith(".xlsx"):
        rows = parse_xlsx(data)
    elif lower.endswith(".csv"):
        rows = parse_csv(data)
    else:
        raise ValueError("Поддерживаются только файлы .xlsx и .csv")
    if not rows:
        raise ValueError("Файл не содержит строк товаров")
    if len(rows) > 5000:
        raise ValueError("За один импорт поддерживается не более 5000 товаров")
    missing = [idx + 2 for idx, row in enumerate(rows) if not row.get("sku") or not row.get("name")]
    if missing:
        preview = ", ".join(map(str, missing[:10]))
        raise ValueError(f"В строках {preview} отсутствует SKU или название")
    return rows
